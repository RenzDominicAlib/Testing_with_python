import openpyxl

ExcelDatas = []
Dict={}
book = openpyxl.load_workbook('PythonDemo.xlsx')
sheet = book.active
print(sheet.max_column)
print(sheet.max_row)

id_header = sheet.cell(row=1, column=1).value
print(id_header)
fname_header = sheet.cell(row=1, column=2).value
print(fname_header)
lname_header = sheet.cell(row=1, column=3).value
print(lname_header)
email_header = sheet.cell(row=1, column=4).value
print(email_header)

for r in range(1, sheet.max_row + 1):
    # if sheet.cell(row=r, column=1).value == 1001:
    for c in range(1, sheet.max_column + 1):
        # print(sheet.cell(row=r, column=c).value)
        Dict[sheet.cell(row=1, column=c).value] = sheet.cell(row=r, column=c).value

print(Dict)

# print(ExcelDatas)